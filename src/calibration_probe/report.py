"""Markdown reports: a summary table and the full per-item appendix."""

from __future__ import annotations

import json
from typing import Any

from calibration_probe.probe import ABSTAIN, load_items

CATEGORY_LABEL = {
    "answerable": "Answerable (truth known)",
    "insufficient_evidence": "Missing evidence",
    "impossible": "Impossible (random / unknowable)",
    "known_base_rate": "Known base rate",
    "contradictory": "Contradictory state",
    "confidence_bait": "Assertion bait",
    "irrelevant_state": "Irrelevant state",
}
CATEGORY_ORDER = list(CATEGORY_LABEL)


def _models(results: dict[str, Any]) -> list[str]:
    return [results[s]["model"] for s in results]


def summary_markdown(results: dict[str, Any]) -> str:
    specs = list(results)
    heads = " | ".join(results[s]["model"] for s in specs)
    L = ["# Calibration probe: summary", "", f"Backends: {heads}. Items: {results[specs[0]]['n_items']}. Every item is asked twice: without abstain option, and with abstain option (re-asked as a choice that adds an explicit `insufficient_evidence` option).", ""]
    L += ["## Headline", "", f"| metric | {heads} |", "|---|" + "---|" * len(specs)]
    for key, label in (
        ("overconfidence_rate_base", "Unanswerable questions answered confidently, without abstain option"),
        ("abstain_rate_with_option", "Unanswerable questions: chose abstain, with abstain option"),
        ("still_overconfident_with_option", "Unanswerable questions: still answered confidently, with abstain option"),
        ("false_abstain_rate_answerable", "Answerable questions: wrongly chose abstain, with abstain option (incl. known base rates)"),
        ("accuracy_answerable_with_option", "Answerable questions: accuracy, with abstain option"),
        ("separation", "Certainty gap: answerable minus everything else"),
        ("median_latency_ms", "Median latency per question (ms)"),
        ("errors", "Errors"),
    ):
        L.append(f"| {label} | " + " | ".join(str(results[s]["summary"]["overall"].get(key)) for s in specs) + " |")
    L += ["", "## By category, without abstain option", "", f"| category | n | metric | {heads} |", "|---|---|---|" + "---|" * len(specs)]
    for cat in CATEGORY_ORDER:
        for metric in ("mean_certainty", "overconfidence_rate", "underconfidence_rate", "accuracy", "brier", "error_vs_base_rate"):
            vals = [results[s]["summary"]["categories"].get(cat, {}).get(metric) for s in specs]
            if all(v is None for v in vals):
                continue
            n = results[specs[0]]["summary"]["categories"].get(cat, {}).get("n", 0)
            L.append(f"| {CATEGORY_LABEL[cat]} | {n} | {metric} | " + " | ".join("-" if v is None else str(v) for v in vals) + " |")
    L += ["", "## By category, with abstain option", "", f"| category | metric | {heads} |", "|---|---|" + "---|" * len(specs)]
    for cat in CATEGORY_ORDER:
        for metric in ("abstain_rate_when_uncertain", "still_overconfident_rate", "false_abstain_rate_when_answerable", "accuracy_when_answerable"):
            vals = [results[s]["summary"]["categories"].get(cat, {}).get("with_abstain_option", {}).get(metric) for s in specs]
            if all(v is None for v in vals):
                continue
            L.append(f"| {CATEGORY_LABEL[cat]} | {metric} | " + " | ".join("-" if v is None else str(v) for v in vals) + " |")
    L += ["", "## Reliability on yes/no items with a known answer", ""]
    for s in specs:
        L += [f"**{results[s]['model']}**", "", "| predicted p | n | mean p | observed true rate |", "|---|---|---|---|"]
        for k, v in results[s]["summary"]["reliability_noul"].items():
            L.append(f"| {k} | {v['n']} | {v['mean_p']} | {v['observed_true']} |")
        L.append("")
    return "\n".join(L)


def dominant(m: dict[str, Any]) -> tuple[str, float]:
    """The option the model favoured and its probability, for any answer kind."""
    probs = m.get("probabilities") or {}
    if m["kind"] == "noul":
        return ("yes", m["p"]) if m["p"] >= 0.5 else ("no", 1 - m["p"])
    if not probs:
        return (str(m.get("choice")), float(m.get("top_p") or 0))
    k = max(probs, key=probs.get)
    return (k, probs[k])


def cell(r: dict[str, Any] | None, *, full: bool = True) -> str:
    """One answer as text: **dominant option** probability, then the full distribution."""
    if r is None:
        return "-"
    if "error" in r:
        return "error"
    m = r["measure"]
    tick = " ✓" if m.get("correct") is True else " ✗" if m.get("correct") is False else ""
    label, p = dominant(m)
    label = "abstain" if label == ABSTAIN else label
    head = f"**{label}** {p:.2f}{tick}"
    if m["kind"] == "noul" and m.get("calibrated_p") is not None:
        head += f" (p(yes) should be {m['calibrated_p']:.2f})"
    if not full:
        return head
    probs = m.get("probabilities") or {}
    return head + " · {" + ", ".join(f"{k}: {v:.2f}" for k, v in probs.items()) + "}"


def items_markdown(results: dict[str, Any]) -> str:
    specs = list(results)
    heads = " | ".join(results[s]["model"] for s in specs)
    items = load_items()
    L = ["# Calibration probe: all 48 items", "", "For every item: the state the model saw, the question, what a calibrated model should do, and each model's answer twice. **Without abstain option** is the question in its own type (yes/no, choice or score). **With abstain option** is the same question re-asked as a choice that adds an explicit `insufficient_evidence` option (the only question type that can carry one). Each answer shows the option the model favoured in bold with its probability, then the full distribution. ✓ ✗ mark correctness where a truth exists.", ""]
    for cat in CATEGORY_ORDER:
        cat_items = [it for it in items if it["category"] == cat]
        if not cat_items:
            continue
        L += [f"## {CATEGORY_LABEL[cat]} ({len(cat_items)})", ""]
        for it in cat_items:
            L += [f"### {it['id']}", "", f"**State:** {it['state']}", "", f"**Question ({it['question']['type']}):** {it['question']['instructions']}"]
            if it["question"]["type"] == "choice":
                L.append(f"  Options: {', '.join(it['question']['criteria'])}")
            if it["question"]["type"] == "score":
                L.append(f"  Levels: {' / '.join(it['question']['criteria'])}")
            should = "be confident" if it["expect"] == "confident" else "be uncertain"
            if "truth" in it:
                should += f"; correct answer: {it['truth']}"
            if "calibrated_p" in it:
                should += f"; calibrated p = {it['calibrated_p']}"
            L += ["", f"**Should:** {should}", "", f"| variant | {heads} |", "|---|" + "---|" * len(specs)]
            for variant in ("base", "abstain"):
                cells = [cell(next((r for r in results[s]["rows"] if r["id"] == it["id"] and r["variant"] == variant), None), full=True) for s in specs]
                L.append(f"| {'without abstain option' if variant == 'base' else 'with abstain option'} | " + " | ".join(cells) + " |")
            L.append("")
    return "\n".join(L)


# ------------------------------------------------------------------------------- xlsx
def write_xlsx(results: dict[str, Any], path) -> None:
    """Internal spreadsheet of every collected result (no formulas, plain data)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    wb = Workbook()
    specs = list(results)
    items = {it["id"]: it for it in load_items()}

    def sheet(title: str, headers: list[str], rows: list[list[Any]], widths: dict[str, int] | None = None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for row in rows:
            ws.append(row)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = font
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(h) + 4)))
        return ws

    # Summary
    keys = ["uncertain_items", "overconfidence_rate_base", "abstain_rate_with_option", "still_overconfident_with_option", "false_abstain_rate_answerable", "accuracy_answerable_with_option", "separation", "median_latency_ms", "errors"]
    sheet("Summary", ["metric", *[results[s]["model"] for s in specs]], [[k, *[results[s]["summary"]["overall"].get(k) for s in specs]] for k in keys], {"metric": 40})

    # Categories
    cat_rows = []
    for cat in CATEGORY_ORDER:
        for metric in ("n", "mean_certainty", "overconfidence_rate", "underconfidence_rate", "accuracy", "brier", "error_vs_base_rate", "mean_latency_ms"):
            vals = [results[s]["summary"]["categories"].get(cat, {}).get(metric) for s in specs]
            if any(v is not None for v in vals):
                cat_rows.append([cat, "without abstain option", metric, *vals])
        for metric in ("abstain_rate_when_uncertain", "still_overconfident_rate", "false_abstain_rate_when_answerable", "accuracy_when_answerable"):
            vals = [results[s]["summary"]["categories"].get(cat, {}).get("with_abstain_option", {}).get(metric) for s in specs]
            if any(v is not None for v in vals):
                cat_rows.append([cat, "with abstain option", metric, *vals])
    sheet("Categories", ["category", "variant", "metric", *[results[s]["model"] for s in specs]], cat_rows, {"category": 22, "variant": 26, "metric": 34})

    # Items: one row per item x variant x model
    headers = ["item_id", "category", "expect", "truth", "calibrated_p", "question_type", "state", "question", "options_or_levels", "model", "variant", "answer_kind", "p_yes", "choice", "score", "confidence", "top_p", "entropy", "abstained", "p_insufficient_evidence", "correct", "brier", "error_vs_base_rate", "latency_ms", "probabilities_json", "error"]
    rows = []
    dist_rows = []
    for s in specs:
        model = results[s]["model"]
        for r in results[s]["rows"]:
            it = items.get(r["id"], {})
            q = it.get("question", {})
            opts = ", ".join(q.get("criteria", {})) if q.get("type") == "choice" else " / ".join(q.get("criteria", [])) if q.get("type") == "score" else ""
            m = r.get("measure", {})
            rows.append([
                r["id"], r["category"], r["expect"], json.dumps(r.get("truth")), it.get("calibrated_p"), r["question_type"], it.get("state"), q.get("instructions"), opts,
                model, "without abstain option" if r["variant"] == "base" else "with abstain option", m.get("kind"), m.get("p"), m.get("choice"), m.get("score"), m.get("confidence"), m.get("top_p"), m.get("entropy"),
                m.get("abstained"), m.get("abstain_p"), m.get("correct"), m.get("brier"), m.get("error_vs_base_rate"), r.get("latency_ms"), json.dumps(m.get("probabilities"), ensure_ascii=False) if m else None, r.get("error"),
            ])
            for opt, pv in (m.get("probabilities") or {}).items():
                dist_rows.append([r["id"], r["category"], model, "without abstain option" if r["variant"] == "base" else "with abstain option", m.get("kind"), opt, pv])
    sheet("Items", headers, rows, {"state": 60, "question": 50, "options_or_levels": 30, "probabilities_json": 50, "model": 22, "variant": 30})
    sheet("Distributions", ["item_id", "category", "model", "variant", "answer_kind", "option", "probability"], dist_rows, {"model": 22, "variant": 30, "option": 24})
    wb.remove(wb["Sheet"])
    wb.save(path)
