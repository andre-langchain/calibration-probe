from __future__ import annotations

import pytest
from langchain_typesafe import ChoiceAnswer, NoulAnswer, ScoreAnswer

from calibration_probe import probe as cp
from calibration_probe.backends import FakeBackend, get_backend
from calibration_probe.report import items_markdown, summary_markdown


def test_items_valid():
    items = cp.load_items()
    assert len(items) == 48
    assert {"answerable", "insufficient_evidence", "impossible", "known_base_rate", "contradictory", "confidence_bait", "irrelevant_state"} == {i["category"] for i in items}


def test_abstain_variants_for_every_type():
    noul = cp.abstain_question({"type": "noul", "instructions": "?"})
    assert set(noul.criteria) == {"yes", "no", cp.ABSTAIN}
    choice = cp.abstain_question({"type": "choice", "instructions": "?", "criteria": {"a": "A"}})
    assert set(choice.criteria) == {"a", cp.ABSTAIN}
    score = cp.abstain_question({"type": "score", "instructions": "?", "criteria": ["lo", "hi"]})
    assert set(score.criteria) == {"level_0", "level_1", cp.ABSTAIN}
    assert cp.abstain_truth({"question": {"type": "noul"}, "truth": False}) == "no"
    assert cp.abstain_truth({"question": {"type": "score"}, "truth": 2}) == "level_2"
    assert cp.abstain_truth({"question": {"type": "choice"}}) is None


def test_measures():
    m = cp.measure(NoulAnswer(type="noul", noul=0.27), calibrated_p=0.5)
    assert m["error_vs_base_rate"] == pytest.approx(0.23) and m["confident"]
    m = cp.measure(ChoiceAnswer(type="choice", choice=cp.ABSTAIN, probabilities={"a": 0.02, cp.ABSTAIN: 0.98}, confidence=0.96), truth="a")
    assert m["abstained"] and m["correct"] is False
    m = cp.measure(ScoreAnswer(type="score", score=0.0, legend={0: "a", 1: "b"}, probabilities={0: 0.99, 1: 0.01}, confidence=0.98), truth=0)
    assert m["correct"] and m["confident"]


def test_end_to_end_with_fake_backend():
    res = cp.run("fake:uniform")
    s = res["summary"]
    assert s["overall"]["errors"] == 0
    # uniform fake never answers confidently, so it is never overconfident and never abstains
    assert s["overall"]["overconfidence_rate_base"] == 0.0
    assert s["overall"]["abstain_rate_with_option"] == 0.0
    md = summary_markdown({"fake:uniform": res})
    assert "Headline" in md
    appendix = items_markdown({"fake:uniform": res})
    assert appendix.count("### ") == 48


def test_backend_registry():
    assert isinstance(get_backend("fake:x"), FakeBackend)
    with pytest.raises(Exception):
        get_backend("nope")


def test_measures_keep_full_distribution_and_xlsx_export(tmp_path):
    m = cp.measure(ChoiceAnswer(type="choice", choice="a", probabilities={"a": 0.7, "b": 0.3}, confidence=0.4))
    assert m["probabilities"] == {"a": 0.7, "b": 0.3}
    m = cp.measure(NoulAnswer(type="noul", noul=0.2))
    assert m["probabilities"] == {"yes": 0.2, "no": 0.8}
    res = cp.run("fake:uniform")
    from calibration_probe.report import write_xlsx
    from openpyxl import load_workbook

    out = tmp_path / "probe.xlsx"
    write_xlsx({"fake:uniform": res}, out)
    wb = load_workbook(out)
    assert set(wb.sheetnames) == {"Summary", "Categories", "Items", "Distributions"}
    items = wb["Items"]
    assert items.max_row == 1 + 48 * 2  # one row per item x variant for one model
    assert wb["Distributions"].max_row > 48 * 2
